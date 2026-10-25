import contextlib
import importlib.util
import io
import json
import os
from pathlib import Path
import shlex
import tempfile
import unittest
from unittest.mock import Mock, patch

import requests
from openpyxl import load_workbook

from execution_report import ExecutionReport, prepare_report, record_result, with_report

ROOT = Path(__file__).resolve().parents[1]
MODULES = ('approve', 'reject', 'clear-assignment', 'clear-petugas', 'assign',
           'change-mode', 'dtsen', 'tarik-data', 'rekap', 'tandai-selesai')
PERIOD = '00000000-0000-0000-0000-000000000001'


def load_module(name):
    spec = importlib.util.spec_from_file_location(name.replace('-', '_') + '_test', ROOT / name / 'hit_endpoint.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@contextlib.contextmanager
def folder():
    previous = Path.cwd()
    with tempfile.TemporaryDirectory() as temporary:
        os.chdir(temporary)
        try:
            yield Path(temporary)
        finally:
            os.chdir(previous)


def read_report(path):
    workbook = load_workbook(path)
    try:
        summary = dict(workbook['Ringkasan'].values)
        values = list(workbook['Hasil'].values)
        return summary, [dict(zip(values[0], row)) for row in values[1:]]
    finally:
        workbook.close()


def response(status=200, success=True):
    result = Mock(status_code=status, headers={'Content-Type': 'application/json'}, text='dummy response')
    result.json.return_value = {'success': success}
    return result


def write_curl(path='curl.txt'):
    payload = {'assignmentExtraParam': {'surveyPeriodId': PERIOD}}
    Path(path).write_text("curl 'https://example.invalid/datatable-all-user-survey-periode' "
                          "-H 'Content-Type: application/json' --data-raw " + shlex.quote(json.dumps(payload)))


class ExcelReportTests(unittest.TestCase):
    def test_literal_cells_and_atomic_updates_and_archives(self):
        with folder() as directory:
            report = ExecutionReport('test', directory)
            report.prepare([{'id': '0317504000600010', 'nama': '=1+2'}], PERIOD)
            report.record('0317504000600010', 'GAGAL', '=HYPERLINK("https://example.invalid")', 400)
            report.record('0317504000600010', 'BERHASIL', 'updated', 200)
            summary, rows = read_report(report.latest)
            self.assertEqual(summary['Total target'], 1)
            self.assertEqual(summary['BERHASIL'], 1)
            self.assertEqual(rows[0]['ID Target'], '0317504000600010')
            self.assertEqual(rows[0]['HTTP'], '200')
            workbook = load_workbook(report.latest)
            try:
                self.assertEqual(workbook['Hasil']['B2'].data_type, 's')
                self.assertEqual(workbook['Hasil']['B2'].value, '=1+2')
            finally:
                workbook.close()
            previous_archive = report.archive.read_bytes()
            ExecutionReport('test', directory)
            self.assertEqual(report.archive.read_bytes(), previous_archive)
            self.assertEqual(len(list((directory / 'reports').glob('*.xlsx'))), 2)
            self.assertFalse(list(directory.glob('.report-*')))

    def test_interrupt_preserves_success_and_unprocessed(self):
        with folder() as directory:
            @with_report('interrupt')
            def run():
                prepare_report([{'id': 'done'}, {'id': 'pending'}], PERIOD)
                record_result('done', 'BERHASIL')
                raise KeyboardInterrupt
            with self.assertRaises(KeyboardInterrupt):
                run()
            summary, rows = read_report(directory / 'laporan_eksekusi_interrupt.xlsx')
            self.assertEqual(summary['Proses'], 'Dihentikan pengguna')
            self.assertEqual([row['Status'] for row in rows], ['BERHASIL', 'BELUM DIPROSES'])

    def test_failure_before_targets_still_writes_report_without_exception_secret(self):
        with folder() as directory:
            @with_report('failure')
            def run():
                raise RuntimeError('secret-auth-url')
            with self.assertRaises(RuntimeError):
                run()
            summary, rows = read_report(directory / 'laporan_eksekusi_failure.xlsx')
            self.assertEqual(summary['Proses'], 'Terhenti')
            self.assertEqual(rows, [])
            self.assertNotIn('secret-auth-url', str(summary))

    def test_all_ten_modules_write_excel_on_preflight_without_network(self):
        for name in MODULES:
            with self.subTest(module=name), folder() as directory:
                module = load_module(name)
                overrides = {'MODULE_DIR': directory} if hasattr(module, 'MODULE_DIR') else {}
                with (patch.multiple(module, **overrides) if overrides else contextlib.nullcontext()), patch('builtins.input', return_value=''), \
                        patch('requests.sessions.Session.request', side_effect=AssertionError('No network')) as network, \
                        patch.dict(os.environ, {'APPROVE_CURL_FILE': str(directory / 'missing.curl')}), \
                        contextlib.redirect_stdout(io.StringIO()):
                    try:
                        module.main()
                    except SystemExit:
                        pass
                network.assert_not_called()
                path = directory / f'laporan_eksekusi_{name.replace("-", "_")}.xlsx'
                self.assertTrue(path.exists())
                self.assertEqual(read_report(path)[1], [])
                self.assertEqual(len(list((directory / 'reports').glob('*.xlsx'))), 1)

    def test_approve_reports_success_business_failure_network_error_expired_and_pending(self):
        module = load_module('approve')
        with folder() as directory:
            write_curl()
            Path('id_spesifik.txt').write_text('ok\nfailed\nerror\nexpired\npending\n')
            replies = [response(), response(400, False), requests.Timeout('dummy timeout'), response(401)]
            with patch.dict(os.environ, {'APPROVE_CURL_FILE': str(directory / 'curl.txt')}), \
                    patch('builtins.input', side_effect=['4', 'y']), \
                    patch.object(module.requests, 'post', side_effect=replies) as post, \
                    patch.object(module.time, 'sleep'), contextlib.redirect_stdout(io.StringIO()):
                module.main()
            self.assertEqual(post.call_count, 4)
            summary, rows = read_report(directory / 'laporan_eksekusi_approve.xlsx')
            self.assertEqual([row['Status'] for row in rows],
                             ['BERHASIL', 'GAGAL', 'ERROR', 'SESI EXPIRED', 'BELUM DIPROSES'])
            self.assertEqual(summary['Total target'], 5)
            self.assertEqual(rows[0]['Periode'], PERIOD)

    def test_approve_cancel_makes_no_requests_and_targets_remain_pending(self):
        module = load_module('approve')
        with folder() as directory:
            write_curl()
            Path('id_spesifik.txt').write_text('cancelled\n')
            with patch.dict(os.environ, {'APPROVE_CURL_FILE': str(directory / 'curl.txt')}), \
                    patch('builtins.input', side_effect=['4', 'n']), \
                    patch.object(module.requests, 'post') as post, contextlib.redirect_stdout(io.StringIO()):
                module.main()
            post.assert_not_called()
            self.assertEqual(read_report(directory / 'laporan_eksekusi_approve.xlsx')[1][0]['Status'],
                             'BELUM DIPROSES')

    def test_reject_reports_success_and_session_expired(self):
        module = load_module('reject')
        with folder() as directory:
            write_curl()
            Path('id_spesifik.txt').write_text('ok\nexpired\npending\n')
            with patch('builtins.input', side_effect=['3', 'y']), \
                    patch.object(module.requests, 'post', side_effect=[response(), response(403)]) as post, \
                    contextlib.redirect_stdout(io.StringIO()):
                module.main()
            self.assertEqual(post.call_count, 2)
            _, rows = read_report(directory / 'laporan_eksekusi_reject.xlsx')
            self.assertEqual([row['Status'] for row in rows], ['BERHASIL', 'SESI EXPIRED', 'BELUM DIPROSES'])


    def test_clear_assignment_reports_each_id_in_successful_batch(self):
        module = load_module('clear-assignment')
        with folder() as directory:
            write_curl('curl_clear.txt')
            Path('ids.json').write_text(json.dumps([{'id': 'a'}, {'id': 'b'}]))
            with patch('builtins.input', return_value='y'), \
                    patch.object(module.requests, 'post', return_value=response()) as post, \
                    contextlib.redirect_stdout(io.StringIO()):
                module.main()
            self.assertEqual(post.call_count, 1)
            _, rows = read_report(directory / 'laporan_eksekusi_clear_assignment.xlsx')
            self.assertEqual([row['ID Target'] for row in rows], ['a', 'b'])
            self.assertEqual([row['Status'] for row in rows], ['BERHASIL', 'BERHASIL'])

    def test_clear_petugas_reports_success_expired_and_remaining(self):
        module = load_module('clear-petugas')
        with folder() as directory:
            write_curl()
            Path('ids.json').write_text(json.dumps([
                {'allocationId': key, 'regionCode': '0317504000600010', 'email': 'dummy@example.invalid'}
                for key in ('ok', 'expired', 'pending')]))
            with patch('builtins.input', return_value='y'), \
                    patch.object(module.requests, 'delete', side_effect=[response(), response(403)]) as delete, \
                    contextlib.redirect_stdout(io.StringIO()):
                module.main()
            self.assertEqual(delete.call_count, 2)
            _, rows = read_report(directory / 'laporan_eksekusi_clear_petugas.xlsx')
            self.assertEqual([row['Status'] for row in rows], ['BERHASIL', 'SESI EXPIRED', 'BELUM DIPROSES'])

    def test_assign_validation_failure_is_in_excel_as_well_as_csv(self):
        module = load_module('assign')
        with folder() as directory:
            Path('curl.txt').touch()
            module.pd.DataFrame([{'idsbr': '0000123', 'email_pencacah': '', 'email_pengawas': ''}]).to_excel('assign.xlsx', index=False)
            with patch.object(module, 'parse_curl', return_value=('url', 'dummy', 'dummy')), \
                    patch.object(module, 'extract_survey_period_id_from_file', return_value=PERIOD), \
                    patch.object(module, 'fetch_all_users', return_value={}), \
                    patch.object(module, 'fetch_all_samples', return_value={}), \
                    patch.object(module, 'fetch_single_sample_on_demand', return_value=None), \
                    patch.object(module, 'fetch_single_user_on_demand', return_value=None), \
                    patch('builtins.input', return_value='y'), \
                    patch('requests.sessions.Session.request', side_effect=AssertionError('No network')) as network, \
                    contextlib.redirect_stdout(io.StringIO()):
                module.main()
            network.assert_not_called()
            _, rows = read_report(directory / 'laporan_eksekusi_assign.xlsx')
            self.assertEqual(rows[0]['ID Target'], '0000123')
            self.assertEqual(rows[0]['Status'], 'GAGAL')
            self.assertTrue(Path('laporan_hasil_assign.csv').exists())

    def test_change_mode_validation_failure_is_in_excel(self):
        module = load_module('change-mode')
        with folder() as directory:
            Path('curl.txt').touch()
            module.pd.DataFrame([{'idsbr': '0000123', 'mode': 'invalid'}]).to_excel('change_mode.xlsx', index=False)
            with patch.object(module, 'parse_curl', return_value=('url', 'dummy', 'dummy', '', '')), \
                    patch.object(module, 'extract_survey_period_id_from_file', return_value=PERIOD), \
                    patch.object(module, 'find_excel_file', return_value='change_mode.xlsx'), \
                    patch.object(module, 'detect_columns', return_value=('idsbr', None, 'mode')), \
                    patch.object(module.time, 'sleep'), patch('builtins.input', return_value='y'), \
                    patch('requests.sessions.Session.request', side_effect=AssertionError('No network')) as network, \
                    contextlib.redirect_stdout(io.StringIO()):
                module.main()
            network.assert_not_called()
            _, rows = read_report(directory / 'laporan_eksekusi_change_mode.xlsx')
            self.assertEqual(rows[0]['Status'], 'GAGAL')
            self.assertTrue(Path('laporan_hasil_change_mode.xlsx').exists())

    def test_dtsen_missing_region_is_reported_and_input_workbook_is_saved(self):
        module = load_module('dtsen')
        with folder() as directory:
            write_curl()
            workbook = module.openpyxl.Workbook()
            workbook.active.append(['column'] * 10)
            workbook.active.append(['', '', '0000', 'dummy', '', '', 'Dummy Name'])
            workbook.save('dtsen_sample.xlsx')
            workbook.close()
            with patch.object(module, 'DEFAULT_INPUT_EXCEL', str(directory / 'dtsen_sample.xlsx')), \
                    patch.object(module, 'resolve_group_id', return_value='dummy'), \
                    patch.object(module, 'load_or_fetch_region_db', return_value={}), \
                    patch.object(module, 'get_region4_id', return_value=None), \
                    patch('requests.sessions.Session.request', side_effect=AssertionError('No network')) as network, \
                    contextlib.redirect_stdout(io.StringIO()):
                module.main()
            network.assert_not_called()
            _, rows = read_report(directory / 'laporan_eksekusi_dtsen.xlsx')
            self.assertEqual(rows[0]['Status'], 'GAGAL')
            self.assertEqual(rows[0]['Nama'], 'Dummy Name')

    def test_rekap_first_request_expired_still_writes_excel(self):
        module = load_module('rekap')
        with folder() as directory:
            parsed = {'survey_period_id': PERIOD, 'xsrf_token': ''}
            people = [{'email': 'dummy@example.invalid', 'nama': 'Dummy', 'role': 'PML'}]
            with patch.object(module, 'parse_curl', return_value=parsed), \
                    patch.object(module, 'build_request_headers', return_value={}), \
                    patch.object(module, 'parse_excel_petugas', return_value=people), \
                    patch.object(module, 'load_progress', return_value={}), \
                    patch.object(module, 'fetch_progress', return_value=(False, {}, 'SESSION_EXPIRED')), \
                    patch.object(module, 'write_output_excel') as output, \
                    contextlib.redirect_stdout(io.StringIO()):
                module.main()
            output.assert_called_once()
            _, rows = read_report(directory / 'laporan_eksekusi_rekap.xlsx')
            self.assertEqual(rows[0]['Status'], 'SESI EXPIRED')

    def test_tarik_data_invalid_target_is_recorded_even_without_valid_targets(self):
        module = load_module('tarik-data')
        with folder() as directory:
            write_curl()
            Path('config.json').write_text(json.dumps({'groupId': PERIOD, 'surveyPeriodId': PERIOD}))
            Path(module.REGION_DB_FILE).write_text(json.dumps({'0000': {'name': 'Dummy'}}))
            module.pd.DataFrame({'subsls_code': ['0000']}).to_excel(module.TARGET_FILE, index=False)
            with patch.object(module, 'resolve_region_hierarchy', return_value=None), \
                    patch('requests.sessions.Session.request', side_effect=AssertionError('No network')) as network, \
                    contextlib.redirect_stdout(io.StringIO()):
                module.main()
            network.assert_not_called()
            _, rows = read_report(directory / 'laporan_eksekusi_tarik_data.xlsx')
            self.assertEqual(rows[0]['ID Target'], '0000')
            self.assertEqual(rows[0]['Status'], 'GAGAL')


if __name__ == '__main__':
    unittest.main()

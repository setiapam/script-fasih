import copy
import datetime
import json
from pathlib import Path
import shlex
import time
from urllib.parse import parse_qs, unquote, urlparse

import requests
from openpyxl import Workbook, load_workbook


MODULE_DIR = Path(__file__).resolve().parent
BASE_URL = "https://fasih-sm.bps.go.id/app/api/assignment-general/api/assignment-region"
REPORT_FILE = 'laporan_hasil_tandai_selesai.xlsx'
REPORT_COLUMNS = ('Periode', 'Kode Sub SLS', 'Status', 'Keterangan', 'Waktu')
PAYLOAD_FIELDS = (
    "id", "surveyPeriodId", "smallestRegionFullCode", "region",
    "regionId", "regionGroupId", "doneListing", "doneTarikSample",
)


class SessionExpiredException(Exception):
    pass


def parse_curl(content):
    """Baca cURL bash sebagai teks, tanpa mengeksekusi perintah shell."""
    tokens = shlex.split(content.replace('\\\r\n', ' ').replace('\\\n', ' '))
    headers = requests.structures.CaseInsensitiveDict()
    url, body = "", {}
    i = 0
    while i < len(tokens):
        token = tokens[i]
        if token in ('-H', '--header', '-b', '--cookie', '--url',
                     '-d', '--data', '--data-raw', '--data-binary'):
            value = tokens[i + 1]
            if token in ('-H', '--header'):
                key, value = value.split(':', 1)
                headers[key.strip()] = value.strip()
            elif token in ('-b', '--cookie'):
                headers['Cookie'] = value
            elif token == '--url':
                url = value
            else:
                body = json.loads(value)
            i += 2
        else:
            if token.startswith('https://'):
                url = token
            i += 1
    parsed = urlparse(url)
    if (parsed.scheme != 'https' or parsed.netloc != 'fasih-sm.bps.go.id'
            or parsed.path != urlparse(BASE_URL).path + '/datatable'):
        raise ValueError("Gunakan cURL pencarian assignment-region/datatable.")
    period = parse_qs(parsed.query).get('periodeId', [''])[0]
    if not period:
        raise ValueError("Parameter periodeId tidak ditemukan pada URL cURL.")
    if not headers.get('X-XSRF-TOKEN'):
        for cookie in headers.get('Cookie', '').split(';'):
            key, _, value = cookie.strip().partition('=')
            if key == 'XSRF-TOKEN':
                headers['X-XSRF-TOKEN'] = unquote(value)
    if not headers.get('Cookie') or not headers.get('X-XSRF-TOKEN'):
        raise ValueError("Cookie atau X-XSRF-TOKEN tidak ditemukan pada cURL.")
    for key in ('Host', 'Content-Length'):
        headers.pop(key, None)
    headers['Content-Type'] = 'application/json'
    if not isinstance(body, dict):
        raise ValueError("Payload pencarian harus berupa objek JSON.")
    return url, dict(headers), period, body


def load_targets(path):
    targets = []
    for line_number, line in enumerate(path.read_text(encoding='utf-8-sig').splitlines(), 1):
        code = line.strip()
        if not code or code.startswith('#'):
            continue
        if not code.isascii() or not code.isdigit() or len(code) != 16:
            raise ValueError(f"list.txt baris {line_number}: kode Sub SLS harus 16 digit.")
        if code not in targets:
            targets.append(code)
    return targets


def read_response(response):
    if (response.status_code in (401, 403) or
            'text/html' in response.headers.get('Content-Type', '').lower()):
        raise SessionExpiredException("Sesi login habis atau server mengembalikan halaman HTML.")
    response.raise_for_status()
    result = response.json()
    if isinstance(result, dict) and result.get('success') is False:
        raise ValueError(f"API menolak request: {result.get('message', 'success=false')}")
    return result


def extract_items(result):
    if isinstance(result, list):
        return result
    if isinstance(result, dict):
        for key in ('searchData', 'data'):
            value = result.get(key)
            if isinstance(value, list):
                return value
            if isinstance(value, dict):
                return extract_items(value)
    raise ValueError("Format respons datatable tidak dikenali.")


def search_region(url, headers, base_payload, code):
    payload = copy.deepcopy(base_payload)
    payload.update(start=0, length=100, search={'value': code, 'regex': True})
    payload.setdefault('order', [{'column': 0, 'dir': 'asc'}])
    matches, seen = {}, set()
    while True:
        result = read_response(requests.post(url, headers=headers, json=payload, timeout=30))
        items = extract_items(result)
        if not items:
            break
        new_ids = {item.get('id') for item in items if isinstance(item, dict)} - seen
        if not new_ids:
            raise ValueError("Pagination API mengulang halaman yang sama.")
        seen.update(new_ids)
        for item in items:
            if isinstance(item, dict) and str(item.get('smallestRegionFullCode', '')) == code:
                if not item.get('id'):
                    raise ValueError("Assignment wilayah tidak memiliki ID.")
                matches[item['id']] = item
        if len(items) < payload['length']:
            break
        payload['start'] += payload['length']
    if len(matches) != 1:
        raise ValueError(f"Ditemukan {len(matches)} assignment dengan kode persis; harus tepat satu.")
    return next(iter(matches.values()))


def build_done_payload(item, period):
    missing = [key for key in PAYLOAD_FIELDS if key not in item]
    if missing:
        raise ValueError(f"Data assignment tidak lengkap: {', '.join(missing)}")
    if item['surveyPeriodId'] != period:
        raise ValueError("Periode assignment berbeda dengan periode pencarian.")
    if not all(item[key] for key in ('id', 'regionId', 'regionGroupId')) or not isinstance(item['region'], dict):
        raise ValueError("Identitas atau objek region tidak valid.")
    if not all(isinstance(item[key], bool) for key in ('doneListing', 'doneTarikSample')):
        raise ValueError("Status doneListing/doneTarikSample harus boolean.")
    # Endpoint /done menerima status saat ini, persis seperti request browser.
    return {key: copy.deepcopy(item[key]) for key in PAYLOAD_FIELDS}


def log_message(message):
    print(message)
    with (MODULE_DIR / 'execution.log').open('a', encoding='utf-8') as stream:
        stream.write(f"{datetime.datetime.now():%Y-%m-%d %H:%M:%S} {message}\n")


def save_report(results):
    """Perbarui hasil per periode/kode tanpa menghapus target lain dari laporan."""
    path = MODULE_DIR / REPORT_FILE
    rows = {}
    if path.exists():
        previous = load_workbook(path, read_only=True)
        try:
            for row in previous.active.iter_rows(min_row=2, values_only=True):
                rows[(row[0], row[1])] = row
        finally:
            previous.close()
    for result in results:
        row = tuple(result.get(column, '') for column in REPORT_COLUMNS)
        rows[(row[0], row[1])] = row
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Hasil Tandai Selesai'
    sheet.append(REPORT_COLUMNS)
    for row in rows.values():
        sheet.append(row)
        # Kode 16 digit harus tetap berupa teks di Excel.
        sheet.cell(sheet.max_row, 2).number_format = '@'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip('ABCDE', (38, 22, 22, 90, 22)):
        sheet.column_dimensions[column].width = width
    temporary = path.with_suffix('.tmp.xlsx')
    workbook.save(temporary)
    workbook.close()
    temporary.replace(path)


def main():
    print('=' * 60)
    print('          TANDAI SELESAI LISTING WILAYAH')
    print('=' * 60)
    try:
        url, headers, period, base_payload = parse_curl(
            (MODULE_DIR / 'curl.txt').read_text(encoding='utf-8'))
        targets = load_targets(MODULE_DIR / 'list.txt')
    except (OSError, ValueError, IndexError) as exc:
        print(f"[!] {exc}\nLihat panduan tandai-selesai/README.md.")
        return
    if not targets:
        print('[!] list.txt kosong. Isi satu kode Sub SLS per baris.')
        return
    print(f'Periode: {period}\nTarget dari list.txt: {len(targets)} wilayah')
    for code in targets:
        print(f' - {code}')
    if input('Tandai listing selesai untuk daftar ini? (y/N): ').strip().lower() != 'y':
        print('[*] Dibatalkan.')
        return
    (MODULE_DIR / 'config.json').write_text(
        json.dumps({'surveyPeriodId': period}, indent=2), encoding='utf-8')
    success = failed = skipped = 0
    results = [dict(zip(REPORT_COLUMNS, (
        period, code, 'Belum diproses', '',
        datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
    ))) for code in targets]
    save_report(results)
    try:
        for index, code in enumerate(targets, 1):
            prefix = f'[{index}/{len(targets)}] {code}'
            report = results[index - 1]
            try:
                item = search_region(url, headers, base_payload, code)
                payload = build_done_payload(item, period)
                if payload['doneListing']:
                    skipped += 1
                    report.update(Status='Sudah selesai', Keterangan='Listing sudah selesai di server.')
                    log_message(f'{prefix} [SKIPPED] Listing sudah selesai di server.')
                    continue
                result = read_response(requests.post(
                    BASE_URL + '/done', headers=headers, json=payload, timeout=30))
                if not isinstance(result, dict) or result.get('success') is not True:
                    raise ValueError('Respons /done tidak mengonfirmasi success=true; periksa status di FASIH.')
                success += 1
                report.update(Status='Berhasil', Keterangan='Listing ditandai selesai.')
                log_message(f'{prefix} [SUKSES] Listing ditandai selesai.')
            except SessionExpiredException as exc:
                report.update(Status='Sesi expired', Keterangan=str(exc))
                raise
            except (requests.RequestException, ValueError) as exc:
                failed += 1
                report.update(Status='Gagal', Keterangan=str(exc))
                log_message(f'{prefix} [GAGAL] {exc}')
            finally:
                report['Waktu'] = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                save_report(results)
            time.sleep(0.3)
    except SessionExpiredException as exc:
        log_message(f'[SESI EXPIRED] {exc} Perbarui tandai-selesai/curl.txt lalu jalankan ulang.')
    except KeyboardInterrupt:
        log_message('[*] Dihentikan pengguna. Jalankan ulang untuk melanjutkan.')
    log_message(f'Ringkasan: berhasil={success}, gagal={failed}, sudah selesai={skipped}, '
                f'belum diproses={len(targets) - success - failed - skipped}, total={len(targets)}')
    print(f'[*] File hasil: {MODULE_DIR / REPORT_FILE}')


if __name__ == '__main__':
    main()

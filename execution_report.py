"""Laporan Excel per eksekusi, tanpa menyimpan sesi atau kredensial."""
from collections import Counter
from contextvars import ContextVar
from datetime import datetime
from functools import wraps
from pathlib import Path
import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

_CURRENT = ContextVar('execution_report', default=None)
COLUMNS = ('ID Target', 'Nama', 'Wilayah', 'Periode', 'Status', 'Keterangan', 'HTTP', 'Waktu')


def _text(value):
    # Sel selalu berupa teks: kode panjang dan pesan berawalan '=' tetap literal.
    return '' if value is None else str(value)


class ExecutionReport:
    def __init__(self, module, directory):
        self.module = module
        self.started = datetime.now().astimezone()
        self.directory = Path(directory).resolve()
        self.latest = self.directory / f'laporan_eksekusi_{module.replace("-", "_")}.xlsx'
        self.archive = self.directory / 'reports' / (
            f'{self.latest.stem}_{self.started.strftime("%Y%m%d_%H%M%S_%f")}.xlsx')
        self.rows = {}
        self.outcome = 'Berjalan'
        self.note = ''
        self.save()

    def prepare(self, targets, period=''):
        for target in targets:
            key = _text(target['id'])
            self.rows[key] = dict(zip(COLUMNS, (
                key, _text(target.get('nama')), _text(target.get('wilayah')),
                _text(period), 'BELUM DIPROSES', '', '', '',
            )))
        self.save()

    def record(self, key, status, message='', http_status=None):
        key = _text(key)
        if key not in self.rows:
            self.prepare([{'id': key}])
        self.rows[key].update(Status=status, Keterangan=_text(message),
                              HTTP=_text(http_status), Waktu=datetime.now().astimezone().isoformat(timespec='seconds'))
        self.save()

    def save(self):
        workbook = Workbook()
        summary = workbook.active
        summary.title = 'Ringkasan'
        counts = Counter(row['Status'] for row in self.rows.values())
        for row in [('Modul', self.module), ('Mulai', self.started.isoformat(timespec='seconds')),
                    ('Diperbarui', datetime.now().astimezone().isoformat(timespec='seconds')),
                    ('Proses', self.outcome), ('Keterangan', self.note), ('Total target', len(self.rows))]:
            summary.append(row)
        for status in ('BERHASIL', 'GAGAL', 'ERROR', 'SESI EXPIRED', 'DILEWATI', 'BELUM DIPROSES'):
            summary.append((status, counts[status]))
        details = workbook.create_sheet('Hasil')
        details.append(COLUMNS)
        for row in self.rows.values():
            details.append([_text(row.get(column)) for column in COLUMNS])
        for sheet in workbook:
            sheet.freeze_panes = 'A2'
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='244062')
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                        cell.number_format = '@'
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(
                    80, max(18, max(len(_text(cell.value)) for cell in column) + 2))
        details.auto_filter.ref = details.dimensions
        try:
            for path in (self.archive, self.latest):
                path.parent.mkdir(parents=True, exist_ok=True)
                fd, temporary = tempfile.mkstemp(prefix='.report-', suffix='.xlsx', dir=path.parent)
                os.close(fd)
                try:
                    workbook.save(temporary)
                    os.replace(temporary, path)
                finally:
                    Path(temporary).unlink(missing_ok=True)
        finally:
            workbook.close()


def with_report(module):
    """Simpan hasil parsial juga ketika return awal, error, atau Ctrl+C."""
    def decorate(function):
        @wraps(function)
        def wrapped(*args, **kwargs):
            directory = function.__globals__.get('MODULE_DIR', Path.cwd())
            report = ExecutionReport(module, directory)
            token = _CURRENT.set(report)
            try:
                return function(*args, **kwargs)
            except KeyboardInterrupt:
                report.outcome = 'Dihentikan pengguna'
                raise
            except BaseException as exc:
                report.outcome = 'Terhenti'
                # Hindari menyalin exception mentah yang bisa memuat URL autentikasi.
                report.note = f'Eksekusi terhenti ({type(exc).__name__}); periksa terminal.'
                raise
            finally:
                try:
                    if report.outcome == 'Berjalan':
                        report.outcome = 'Proses berakhir'
                    report.save()
                    print(f'[*] Laporan Excel: {report.latest}')
                    print(f'[*] Riwayat eksekusi: {report.archive}')
                finally:
                    _CURRENT.reset(token)
        return wrapped
    return decorate


def prepare_report(targets, period=''):
    report = _CURRENT.get()
    if report is not None:
        report.prepare(targets, period)


def record_result(key, status, message='', http_status=None):
    report = _CURRENT.get()
    if report is not None:
        report.record(key, status, message, http_status)


class ReportResults(list):
    """Pertahankan daftar hasil modul sambil menyimpan tiap append ke Excel."""
    def append(self, result):
        super().append(result)
        message = str(result.get('Status', ''))
        lower = message.lower()
        status = ('DILEWATI' if 'skipped' in lower else
                  'BERHASIL' if lower.startswith('berhasil') else
                  'ERROR' if lower.startswith('error') else 'GAGAL')
        details = [f'{key}: {value}' for key, value in result.items()
                   if key not in ('IDSBR', 'Status') and value is not None]
        record_result(result['IDSBR'], status, '; '.join([message] + details))
